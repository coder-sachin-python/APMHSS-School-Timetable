"""Optional Word/PDF/Excel export (install python-docx reportlab openpyxl)."""
from collections import Counter
from timetable_core import SCHOOL, PERIODS, group, school_days


def table_data(context, slots, store):
    year, month, week, grade, section = context
    days, periods = store.layout(grade, section)
    heading=['Day']+[f'P{i+1}\n{a}-{b}' for i,(a,b) in enumerate(periods)]
    lines=[heading]
    for day in days:
        row=[day]
        for p in range(len(periods)):
            slot=slots.get((day,p))
            if not slot: row.append('---');continue
            teacher=slot.teacher or store.teacher_for(grade,section,slot.subject)
            row.append(slot.subject+(f'\n{teacher}' if teacher else ''))
        lines.append(row)
    return lines


def export_word(path,context,slots,store):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT
    data=table_data(context,slots,store)
    year,month,week,grade,section=context
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.55)
    sec.left_margin=sec.right_margin=Inches(.5)
    sec.page_width=Inches(16) if len(data[0])>8 else Inches(13)
    sec.page_height=Inches(8.5)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('OM SAKTHI\n'+SCHOOL+'\nMELMARUVATHUR');r.bold=True;r.font.size=Pt(13)
    p=doc.add_paragraph(f'REVISION TIME TABLE | CLASS {grade}-{section} | {week}, {month} | {year}')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f'Students: {store.students(grade,section)}')
    table=doc.add_table(rows=0, cols=len(data[0]));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(data):
        cells=table.add_row().cells
        for j,item in enumerate(row):
            cells[j].text=str(item)
            for paragraph in cells[j].paragraphs:
                for run in paragraph.runs:
                    run.font.size=Pt(8)
                    if i==0: run.bold=True
    doc.save(path)


def export_excel(path,context,slots,store):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    data=table_data(context,slots,store)
    year,month,week,grade,section=context
    wb=Workbook();ws=wb.active;ws.title=f'{grade}-{section}'
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(data[0]))
    ws.cell(1,1,f'OM SAKTHI | {SCHOOL}')
    ws.cell(1,1).font=Font(size=14,bold=True,color='FFFFFF')
    ws.cell(1,1).fill=PatternFill('solid',fgColor='7A0019')
    ws.cell(2,1,f'Class {grade}-{section} | {week} {month} {year} | Students: {store.students(grade,section)}')
    for r,row in enumerate(data,4):
        for c,value in enumerate(row,1):
            cell=ws.cell(r,c,value)
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border=Border(bottom=Side(style='thin',color='D8DEE8'))
            if r==4:
                cell.fill=PatternFill('solid',fgColor='26394F')
                cell.font=Font(color='FFFFFF',bold=True)
            elif r%2==1:
                cell.fill=PatternFill('solid',fgColor='F0F5FA')
        ws.row_dimensions[r].height=56 if r>4 else 36
    for c in range(1,len(data[0])+1):
        ws.column_dimensions[get_column_letter(c)].width=23 if c>1 else 16
    ws.freeze_panes='B5'
    ws.sheet_view.showGridLines=False
    ws.print_options.horizontalCentered=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.fitToWidth=1
    wb.save(path)


def export_pdf(path,context,slots,store):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape,A3
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,KeepTogether
    from reportlab.lib.units import mm
    from xml.sax.saxutils import escape
    year,month,week,grade,section=context
    matrix=[]
    style=ParagraphStyle('tiny',fontName='Helvetica',fontSize=7.5,leading=10,alignment=TA_CENTER)
    for row in table_data(context,slots,store):
        matrix.append([Paragraph(escape(str(x)).replace('\n','<br/>'),style) for x in row])
    paper=landscape(A3)
    page_width=paper[0]-30*mm
    n=len(matrix[0]); widths=[28*mm]+[(page_width-28*mm)/(n-1)]*(n-1)
    doc=SimpleDocTemplate(path,pagesize=paper,rightMargin=15*mm,leftMargin=15*mm,topMargin=15*mm,bottomMargin=15*mm)
    main_style=ParagraphStyle('title',fontName='Helvetica-Bold',fontSize=13,leading=19,alignment=TA_CENTER)
    title=Paragraph(escape(SCHOOL),main_style)
    info=Paragraph(escape(f'OM SAKTHI | MELMARUVATHUR | Class {grade}-{section} | {week}, {month} {year} | Students: {store.students(grade,section)}'),style)
    t=Table(matrix,colWidths=widths,repeatRows=1,rowHeights=[19*mm]+[28*mm]*(len(matrix)-1))
    t.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#7A0019')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.whitesmoke),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EDF3F8')]),
        ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CBD5E1')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),5),
        ('RIGHTPADDING',(0,0),(-1,-1),5),
    ]))
    doc.build([title,Spacer(1,9),info,Spacer(1,14),t])

EXPORTERS={'Word (.docx)':('.docx',export_word),'PDF (.pdf)':('.pdf',export_pdf),'Excel (.xlsx)':('.xlsx',export_excel)}
