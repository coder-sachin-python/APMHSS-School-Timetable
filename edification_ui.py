"""Tkinter morning-test planner with school holiday calendar and scoped access."""
from __future__ import annotations
import calendar
import csv
from datetime import date, timedelta
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
from timetable_core import GRADES, SCHOOL, default_subjects, section_display, TimetableError
from edification_core import CYCLES, TEST_TIME, parse_date

MAROON='#7A0019'


def show_edification(app):
    root=app.root
    app._maximize_main_window()
    app.reset_root()
    app.mode='edification'
    app.store=app.auth_store
    engine=app.edification
    admin=app.role=='admin'
    editable=app.role in ('admin','editor')
    grades=[g for g in app.allowed_grades() if g in ('6','7','8','9','10','11','12')]

    header=tk.Frame(root,bg=MAROON,padx=14,pady=11)
    header.pack(fill='x')
    tk.Label(header,text=SCHOOL+'   |   EDIFICATION MORNING TEST TIMETABLE',
             bg=MAROON,fg='white',font=('Segoe UI',14,'bold')).pack(side='left')
    ttk.Button(header,text='Choose Timetable',command=app.back_to_mode_selection).pack(side='right',padx=5)
    ttk.Button(header,text='Logout',command=app.logout).pack(side='right',padx=5)
    tk.Label(header,text=f'{app.user} ({app.role})',bg=MAROON,fg='white').pack(side='right',padx=10)
    if not grades:
        ttk.Label(root,text='No Edification subjects assigned. Contact the administrator.',
                  font=('Segoe UI',14)).pack(pady=35)
        return

    top=ttk.Frame(root,padding=(16,12));top.pack(fill='x')
    grade=tk.StringVar(value='12' if '12' in grades else grades[0])
    sections=app.allowed_sections(grade.get())
    section=tk.StringVar(value=sections[0])
    cycle=tk.StringVar(value=CYCLES[0])
    start=tk.StringVar(value=date.today().isoformat())
    selected_day=tk.StringVar(value='')
    status=tk.StringVar(value='Choose the first test date and subjects. Tests run 8:00–8:45 AM.')
    month_date=[date.today().replace(day=1)]
    subjects=[]

    for col,label,var,values,width in [
        (0,'Class',grade,grades,10),(1,'Section',section,sections,18),
        (2,'Test Cycle',cycle,CYCLES,18)]:
        ttk.Label(top,text=label).grid(row=0,column=col,sticky='w',padx=6)
        cb=ttk.Combobox(top,textvariable=var,values=values,state='readonly',width=width)
        cb.grid(row=1,column=col,sticky='ew',padx=6,pady=4)
        if col==0:grade_cb=cb
        elif col==1:section_cb=cb
        else:cycle_cb=cb
    ttk.Label(top,text='First Test Date (YYYY-MM-DD)').grid(row=0,column=3,sticky='w',padx=6)
    ttk.Entry(top,textvariable=start,width=18).grid(row=1,column=3,sticky='ew',padx=6)
    day_label=ttk.Label(top,text='')
    day_label.grid(row=1,column=4,sticky='w',padx=6)
    ttk.Label(top,text=f'MORNING TEST: {TEST_TIME}',font=('Segoe UI',11,'bold'),
              foreground=MAROON).grid(row=0,column=4,columnspan=2,sticky='w',padx=6)
    for c in (1,2,3):top.grid_columnconfigure(c,weight=1)
    ttk.Label(top,textvariable=status,foreground='#325B79',wraplength=1150).grid(
        row=2,column=0,columnspan=6,sticky='w',pady=(6,0))

    permissions_bar=ttk.Frame(root,padding=(12,0));permissions_bar.pack(fill='x')
    ttk.Button(permissions_bar,text='Teacher Directory',command=app.teacher_directory).pack(side='left',padx=3)
    if not admin:
        for label,action in [('Unit Dashboard',app.view_assigned_unit_dashboard),
                             ('Check Test Clashes',lambda:check_test_clashes()),
                             ('Accounts (Admin)',app.manage_users),
                             ('Login History (Admin)',app.login_history),
                             ('Backup (Admin)',app.backup)]:
            ttk.Button(permissions_bar,text=label,command=action).pack(side='left',padx=3)
    body=ttk.Frame(root,padding=(12,4));body.pack(fill='both',expand=True)
    body.grid_columnconfigure(0,weight=0)
    body.grid_columnconfigure(1,weight=1)
    body.grid_rowconfigure(0,weight=1)
    left=ttk.Frame(body,padding=6);left.grid(row=0,column=0,sticky='nsew')
    right=ttk.Frame(body,padding=6);right.grid(row=0,column=1,sticky='nsew')
    right.grid_columnconfigure(0,weight=1);right.grid_rowconfigure(1,weight=1)

    calendar_box=ttk.LabelFrame(left,text='SCHOOL HOLIDAY CALENDAR',padding=8)
    calendar_box.pack(fill='x',pady=4)
    calendar_header=ttk.Frame(calendar_box);calendar_header.pack(fill='x')
    ttk.Button(calendar_header,text='◀',width=3,command=lambda:change_month(-1)).pack(side='left')
    month_label=ttk.Label(calendar_header,text='',font=('Segoe UI',12,'bold'),width=23,anchor='center')
    month_label.pack(side='left',expand=True)
    ttk.Button(calendar_header,text='▶',width=3,command=lambda:change_month(1)).pack(side='right')
    calendar_days=ttk.Frame(calendar_box);calendar_days.pack(pady=5)
    ttk.Label(calendar_box,text='Red = recorded holiday  •  Grey = Sunday\nClick a date to select the test start date.',
              font=('Segoe UI',9),foreground='#53657B').pack(anchor='w')
    holiday_buttons=ttk.Frame(calendar_box);holiday_buttons.pack(fill='x',pady=5)

    subject_frame=ttk.LabelFrame(left,text='MAIN SUBJECTS FOR THIS CYCLE',padding=8)
    subject_frame.pack(fill='both',expand=True,pady=6)
    subject_list=tk.Listbox(subject_frame,height=8,width=38,selectmode='extended',exportselection=False)
    subject_list.pack(fill='both',expand=True)
    subject_input=ttk.Entry(subject_frame,width=30)
    subject_input.pack(fill='x',pady=(6,2))
    subject_buttons=ttk.Frame(subject_frame);subject_buttons.pack(fill='x',pady=3)

    ttk.Label(right,text='TEST DATES • DAYS • SUBJECTS',font=('Segoe UI',14,'bold')).grid(
        row=0,column=0,sticky='w',pady=(5,10))
    tree=ttk.Treeview(right,columns=('No','Date','Day','Time','Subject','Status'),show='headings',height=13)
    for name,width in [('No',48),('Date',120),('Day',120),('Time',130),('Subject',220),('Status',110)]:
        tree.heading(name,text=name);tree.column(name,width=width,anchor='center' if name!='Subject' else 'w')
    tree.grid(row=1,column=0,sticky='nsew')
    scrollbar=ttk.Scrollbar(right,orient='vertical',command=tree.yview)
    scrollbar.grid(row=1,column=1,sticky='ns');tree.configure(yscrollcommand=scrollbar.set)
    commands=ttk.Frame(right);commands.grid(row=2,column=0,sticky='ew',pady=10)
    action_buttons=ttk.Frame(commands);action_buttons.pack(fill='x',pady=(0,5))
    download_buttons=ttk.Frame(commands);download_buttons.pack(fill='x')
    ttk.Label(right,text='Holiday changes do not silently overwrite existing tests.\n'
        'To change the test dates, edit a row or remove and add the cycle again.',
        foreground='#53657B',justify='left').grid(row=3,column=0,sticky='w',pady=8)

    def check_test_clashes():
        """Read-only morning-test date check, scoped to visible subjects."""
        permitted = None if admin else {s.strip().casefold() for s in
                                         app.granted_subjects(grade.get(),section.get())}
        entries = engine.db.execute('''SELECT day,subject,cycle FROM edification_tests
            WHERE grade=? AND section=? ORDER BY day''', (grade.get(),section.get())).fetchall()
        visible = [row for row in entries if permitted is None or
                   row['subject'].strip().casefold() in permitted]
        from collections import Counter
        counts = Counter(row['day'] for row in entries)
        duplicates = sum(1 for row in visible if counts[row['day']] > 1)
        holidays = engine.holidays()
        holiday_hits = sum(1 for row in visible if row['day'] in holidays or
                           parse_date(row['day']).weekday() == 6)
        messagebox.showinfo('Edification • Date Clash Check',
            f'Your visible morning tests: {len(visible)}\n'
            f'Date conflicts with another test for this section: {duplicates}\n'
            f'Tests on Sundays or recorded holidays: {holiday_hits}\n\n'
            'Read-only check. No tests were changed.', parent=root)

    def chosen_subjects():return list(subject_list.get(0,tk.END))

    def refresh_subjects(*_):
        subject_list.delete(0,tk.END)
        if not admin:
            options=sorted(app.granted_subjects(grade.get(),section.get()))
        else:
            options=engine.subject_bank(grade.get(),section.get())
        for sub in options:subject_list.insert(tk.END,sub)

    def refresh_tests(*_):
        tree.delete(*tree.get_children())
        for row in engine.tests(grade.get(),section.get(),cycle.get()):
            if not admin and row['subject'].strip().casefold() not in {
                subject.strip().casefold() for subject in app.granted_subjects(grade.get(),section.get())}:
                continue
            d=parse_date(row['day'])
            holiday=engine.holidays(d.year,d.month).get(row['day'])
            tree.insert('',tk.END,values=(row['position'],row['day'],d.strftime('%A'),
                        TEST_TIME,row['subject'],'HOLIDAY!' if holiday else 'Scheduled'))
        status.set(f'{cycle.get()} • Class {grade.get()}-{section_display(grade.get(),section.get())} • '
                   f'{len(tree.get_children())} visible test(s)')

    def refresh_calendar():
        for child in calendar_days.winfo_children():child.destroy()
        current=month_date[0]
        month_label.config(text=current.strftime('%B %Y'))
        holidays=engine.holidays(current.year,current.month)
        for col,name in enumerate(('Mon','Tue','Wed','Thu','Fri','Sat','Sun')):
            ttk.Label(calendar_days,text=name,width=5,anchor='center').grid(row=0,column=col,padx=2,pady=2)
        for row,week in enumerate(calendar.monthcalendar(current.year,current.month),1):
            for col,num in enumerate(week):
                if not num:continue
                d=date(current.year,current.month,num)
                is_holiday=d.isoformat() in holidays
                selected=d.isoformat()==start.get()
                bg='#F5C4C4' if is_holiday else '#A5D5FF' if selected else '#E4E7ED' if col==6 else '#EAF4FC'
                tk.Button(calendar_days,text=str(num),width=4,relief='solid' if selected else 'flat',bg=bg,
                          command=lambda x=d:choose_day(x)).grid(row=row,column=col,padx=2,pady=2)

    def change_month(delta):
        old=month_date[0]
        year=old.year+(old.month-1+delta)//12
        month=(old.month-1+delta)%12+1
        month_date[0]=date(year,month,1)
        refresh_calendar()

    def choose_day(d):
        start.set(d.isoformat())
        selected_day.set(d.isoformat())
        reason=engine.holidays(d.year,d.month).get(d.isoformat())
        status.set(f'{d:%A, %d %B %Y}' + (f' • HOLIDAY: {reason}' if reason else ''))

    def add_holiday():
        if not app.require_admin():return
        try:d=parse_date(start.get())
        except TimetableError as err:messagebox.showerror('Date',str(err));return
        reason=simpledialog.askstring('School Holiday',f'Reason for {d:%A, %d %B %Y}:',parent=root)
        if reason is None:return
        try:engine.set_holiday(d.isoformat(),reason)
        except TimetableError as err:messagebox.showerror('Holiday',str(err));return
        refresh_calendar();refresh_tests()

    def remove_holiday():
        if not app.require_admin():return
        try:engine.delete_holiday(start.get())
        except TimetableError as err:messagebox.showerror('Date',str(err));return
        refresh_calendar();refresh_tests()

    if admin:
        ttk.Button(holiday_buttons,text='Add Holiday',command=add_holiday).pack(side='left',padx=2)
        ttk.Button(holiday_buttons,text='Remove Holiday',command=remove_holiday).pack(side='left',padx=2)

    def add_subject():
        if not app.require_admin():return
        subject=subject_input.get().strip()
        if not subject or len(subject)>65:
            messagebox.showwarning('Subject','Enter a subject name (up to 65 characters).');return
        if subject.casefold() in [s.casefold() for s in chosen_subjects()]:
            messagebox.showwarning('Subject','This subject is already listed.');return
        subject_list.insert(tk.END,subject);subject_input.delete(0,tk.END)

    def remove_subject():
        if not app.require_admin():return
        for i in reversed(subject_list.curselection()):subject_list.delete(i)

    def save_subject_bank():
        if not app.require_admin():return
        try:engine.save_subject_bank(grade.get(),section.get(),chosen_subjects())
        except TimetableError as err:messagebox.showerror('Subjects',str(err),parent=root);return
        messagebox.showinfo('Subjects saved',
            'This main-subject list is saved for the selected class and section.',parent=root)

    if admin:
        ttk.Button(subject_buttons,text='Add Subject',command=add_subject).pack(side='left',padx=3)
        ttk.Button(subject_buttons,text='Remove Selected',command=remove_subject).pack(side='left',padx=3)
        ttk.Button(subject_buttons,text='Save Subject List',command=save_subject_bank).pack(side='left',padx=3)
        ttk.Button(subject_frame,text='Accounts / Teacher Subject Access',command=app.manage_users).pack(
            fill='x',pady=6)

    def generate():
        if not app.require_admin():return
        subjects=chosen_subjects()
        existing=engine.tests(grade.get(),section.get(),cycle.get())
        if existing and not messagebox.askyesno('Regenerate Tests',
            'Replace this cycle and its manual edits? Other cycles are preserved.',parent=root):return
        try:
            engine.generate(grade.get(),section.get(),cycle.get(),start.get(),subjects,replace=True)
            engine.save_subject_bank(grade.get(),section.get(),subjects)
        except TimetableError as err:messagebox.showerror('Cannot generate',str(err),parent=root);return
        refresh_tests()
        app.auth_store.log(app.user,f'Generated {cycle.get()} for {grade.get()}-{section.get()}')

    def add_timetable():
        if not app.require_admin():return
        existing=engine.tests(grade.get(),section.get(),cycle.get())
        if existing:
            messagebox.showinfo('Timetable already exists',
                f'{cycle.get()} already has {len(existing)} tests.\n'
                'Choose another Edification cycle, or remove this timetable first.\n'
                'You can also use Shuffle Subjects to change its order.', parent=root)
            return
        subjects=chosen_subjects()
        if not subjects:
            messagebox.showwarning('Subjects', 'Add main subjects before creating a timetable.', parent=root)
            return
        try:
            engine.generate(grade.get(),section.get(),cycle.get(),start.get(),subjects)
            engine.save_subject_bank(grade.get(),section.get(),subjects)
        except TimetableError as err:
            messagebox.showerror('Cannot add timetable',str(err),parent=root);return
        refresh_tests()
        app.auth_store.log(app.user,f'Added {cycle.get()} for {grade.get()}-{section.get()}')

    def remove_timetable():
        if not app.require_admin():return
        existing=engine.tests(grade.get(),section.get(),cycle.get())
        if not existing:
            messagebox.showinfo('No timetable', 'There is no timetable to remove for this cycle.', parent=root)
            return
        if not messagebox.askyesno('Remove Timetable',
            f'Delete ALL {len(existing)} tests in {cycle.get()} for Class '
            f'{grade.get()}-{section.get()}?\n\nOther cycles and subject lists will be kept.',
            parent=root):return
        engine.delete_cycle(grade.get(),section.get(),cycle.get())
        refresh_tests()
        app.auth_store.log(app.user,f'Removed {cycle.get()} for {grade.get()}-{section.get()}')

    def shuffle_timetable():
        if not app.require_admin():return
        if not messagebox.askyesno('Vary Subject Order',
            'Shuffle the subjects in this timetable?\n\n'
            'Dates, days and test time will stay the same.', parent=root):return
        try:engine.shuffle_cycle(grade.get(),section.get(),cycle.get())
        except TimetableError as err:messagebox.showerror('Shuffle failed',str(err),parent=root);return
        refresh_tests()
        app.auth_store.log(app.user,f'Shuffled {cycle.get()} for {grade.get()}-{section.get()}')

    def swap_subject_order():
        if not app.require_admin():return
        selection=tree.selection()
        if not selection:
            messagebox.showinfo('Select Test','Select a test row to swap its subject.',parent=root);return
        current=int(tree.item(selection[0],'values')[0])
        other=simpledialog.askinteger('Swap Subject Order',
            f'Selected test number: {current}\n\n'
            'Enter the other test number to swap subjects with:',
            parent=root,minvalue=1,maxvalue=len(engine.tests(grade.get(),section.get(),cycle.get())))
        if other is None:return
        try:engine.swap_cycle_subjects(grade.get(),section.get(),cycle.get(),current,other)
        except TimetableError as err:messagebox.showerror('Swap failed',str(err),parent=root);return
        refresh_tests()
        app.auth_store.log(app.user,f'Swapped test subjects in {cycle.get()}')

    def edit_selected():
        if not editable:return
        selected=tree.selection()
        if not selected:
            messagebox.showinfo('Select Test','Select a test row first.');return
        pos,old_date,_,_,old_subject,_=tree.item(selected[0],'values')
        if not admin and old_subject.strip().casefold() not in {
                subject.strip().casefold() for subject in app.granted_subjects(grade.get(),section.get())}:
            messagebox.showerror('Restricted','This subject is not assigned to you.');return
        win=tk.Toplevel(root);win.title('Edit Morning Test');win.geometry('390x250');win.transient(root);win.grab_set()
        form=ttk.Frame(win,padding=20);form.pack(fill='both',expand=True)
        ttk.Label(form,text='Test Date (YYYY-MM-DD)').pack(anchor='w')
        day_var=tk.StringVar(value=old_date)
        ttk.Entry(form,textvariable=day_var).pack(fill='x',pady=5)
        ttk.Label(form,text='Subject').pack(anchor='w')
        sub_var=tk.StringVar(value=old_subject)
        sub_entry=ttk.Combobox(form,textvariable=sub_var,values=chosen_subjects())
        sub_entry.pack(fill='x',pady=5)
        if not admin:sub_entry.configure(state='disabled')
        def apply():
            if not admin and sub_var.get()!=old_subject:
                messagebox.showerror('Restricted','You may edit only your assigned subject.',parent=win);return
            try:engine.update_test(grade.get(),section.get(),cycle.get(),int(pos),day_var.get(),sub_var.get())
            except TimetableError as err:messagebox.showerror('Update failed',str(err),parent=win);return
            win.destroy();refresh_tests()
        ttk.Button(form,text='Save Test Changes',command=apply).pack(fill='x',pady=12)

    def export_csv():
        path=filedialog.asksaveasfilename(defaultextension='.csv',filetypes=[('CSV files','*.csv')],
                                           initialfile=f'Edification_{grade.get()}_{section.get()}.csv')
        if not path:return
        try:
            with open(path,'w',newline='',encoding='utf-8-sig') as f:
                writer=csv.writer(f)
                writer.writerow(['Class','Section','Cycle','Date','Day','Time','Subject'])
                for item in tree.get_children():
                    _,day,weekday,time,subject,_=tree.item(item,'values')
                    writer.writerow([grade.get(),section_display(grade.get(),section.get()),
                                     cycle.get(),day,weekday,time,subject])
        except OSError as err:messagebox.showerror('Export failed',str(err));return
        messagebox.showinfo('Exported','Saved visible test rows to CSV.\n'+path)

    def export_document(kind):
        # Export only rows currently visible to the logged-in user.
        records=[tree.item(item,'values') for item in tree.get_children()]
        if not records:
            messagebox.showwarning('No tests','No authorized tests to download.',parent=root);return
        ext={'Word':'.docx','PDF':'.pdf','Excel':'.xlsx'}[kind]
        path=filedialog.asksaveasfilename(defaultextension=ext,
             initialfile=f'Edification_{grade.get()}_{section.get()}_{cycle.get().replace(" ","")}{ext}',
             filetypes=[(kind,'*'+ext)])
        if not path:return
        columns=['No','Date','Day','Time','Subject','Status']
        try:
            if kind=='Word':
                from docx import Document
                doc=Document();doc.add_heading('APMHSS • Edification Morning Tests',0)
                doc.add_paragraph(f'Class {grade.get()}-{section_display(grade.get(),section.get())} • {cycle.get()}')
                table=doc.add_table(rows=1,cols=len(columns));table.style='Table Grid'
                for i,name in enumerate(columns):table.rows[0].cells[i].text=name
                for row in records:
                    for i,value in enumerate(row):table.add_row().cells[i].text=str(value)
                doc.save(path)
            elif kind=='Excel':
                from openpyxl import Workbook
                from openpyxl.styles import Font,PatternFill
                wb=Workbook();ws=wb.active;ws.title='Edification Tests'
                ws.append(columns)
                for row in records:ws.append(list(row))
                for cell in ws[1]:
                    cell.font=Font(bold=True,color='FFFFFF')
                    cell.fill=PatternFill('solid',fgColor='7A0019')
                for col,width in {'A':9,'B':17,'C':17,'D':20,'E':32,'F':17}.items():
                    ws.column_dimensions[col].width=width
                wb.save(path)
            else:
                from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import landscape,A4
                doc=SimpleDocTemplate(path,pagesize=landscape(A4))
                styles=getSampleStyleSheet()
                rows=[columns]+[[str(v) for v in row] for row in records]
                table=Table(rows,repeatRows=1,colWidths=[45,92,92,100,165,90])
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#7A0019')),
                      ('TEXTCOLOR',(0,0),(-1,0),colors.white),
                      ('GRID',(0,0),(-1,-1),.3,colors.grey),
                      ('FONTSIZE',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
                doc.build([Paragraph('APMHSS • Edification Morning Tests',styles['Title']),Spacer(1,12),table])
        except ImportError as error:
            messagebox.showerror('Missing package',f'{error}\nRun RUN_UPDATED_APP.bat to install packages.',parent=root);return
        except Exception as error:
            messagebox.showerror('Export failed',str(error),parent=root);return
        messagebox.showinfo('Downloaded',f'Saved authorized test rows to:\n{path}',parent=root)

    # Double-click an authorized test to edit it (admin/editor only).
    if editable:
        tree.bind('<Double-1>',lambda event:edit_selected() if tree.identify_row(event.y) else None)
    else:
        tree.bind('<Double-1>',lambda event:messagebox.showinfo('Read-only',
                  'Viewer accounts can view and download assigned tests, but cannot edit them.',parent=root))

    if admin:
        ttk.Button(action_buttons,text='Add Timetable',command=add_timetable).pack(side='left',padx=3)
        ttk.Button(action_buttons,text='Remove Timetable',command=remove_timetable).pack(side='left',padx=3)
        ttk.Button(action_buttons,text='Shuffle Subjects',command=shuffle_timetable).pack(side='left',padx=3)
        ttk.Button(action_buttons,text='Swap Order',command=swap_subject_order).pack(side='left',padx=3)
    if editable:
        ttk.Button(action_buttons,text='Edit Timetable',command=edit_selected).pack(side='left',padx=3)
    for kind in ('Word','PDF','Excel'):
        ttk.Button(download_buttons,text='Download '+kind,command=lambda k=kind:export_document(k)).pack(side='left',padx=2)
    ttk.Button(download_buttons,text='Export Visible Tests (CSV)',command=export_csv).pack(side='left',padx=4)
    ttk.Button(download_buttons,text='Refresh',command=refresh_tests).pack(side='left',padx=4)

    def on_grade(*_):
        options=app.allowed_sections(grade.get())
        section_cb.configure(values=options)
        if section.get() not in options:section.set(options[0] if options else '')
        refresh_subjects();refresh_tests()

    def on_section(*_):refresh_subjects();refresh_tests()
    def on_cycle(*_):
        # Suggest the next available date after the preceding cycle, but allow manual override.
        index=CYCLES.index(cycle.get())
        if index:
            previous=engine.tests(grade.get(),section.get(),CYCLES[index-1])
            if previous:
                last=max(parse_date(row['day']) for row in previous)
                start.set(engine.next_day(last+timedelta(days=1)).isoformat())
        refresh_tests()

    def show_day(*_):
        try:
            day=parse_date(start.get())
            holiday=engine.holidays(day.year,day.month).get(day.isoformat())
            day_label.configure(text=day.strftime('%A')+(' • HOLIDAY' if holiday else ''))
            if (month_date[0].year,month_date[0].month)!=(day.year,day.month):
                month_date[0]=day.replace(day=1)
            refresh_calendar()
        except TimetableError:
            day_label.configure(text='Invalid date')
    start.trace_add('write',show_day)
    grade_cb.bind('<<ComboboxSelected>>',on_grade)
    section_cb.bind('<<ComboboxSelected>>',on_section)
    cycle_cb.bind('<<ComboboxSelected>>',on_cycle)
    refresh_subjects();refresh_calendar();refresh_tests();show_day()
